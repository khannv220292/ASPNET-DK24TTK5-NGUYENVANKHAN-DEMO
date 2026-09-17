# -*- coding: utf-8 -*-
"""Gắn đúng ảnh Chương 4–5 vào báo cáo."""
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm
from docx.text.paragraph import Paragraph

SHOT = Path(r"c:\Users\nguye\AppData\Local\Temp\cursor\screenshots")
OUT = Path(r"C:\Users\nguye\OneDrive\Máy tính\Đồ án")
FIGS = Path(r"C:\Users\nguye\AppData\Local\Temp\baocao_figs")
DOC = OUT / "BaoCao_laptop.khannv.vn_TheoMauChinh.docx"
TMP = OUT / "_tmp_baocao_figs.docx"

# Ưu tiên ảnh có thanh địa chỉ (đã chèn trước đó)
h18 = FIGS / "img_560_0_image19.png"
if not h18.exists():
    h18 = SHOT / "Hinh18_home.png"


def vstack(paths, dest, labels=None):
    imgs = [Image.open(p).convert("RGB") for p in paths]
    w = max(i.width for i in imgs)
    gap = 36
    font = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 28)
    heights = []
    resized = []
    for i, im in enumerate(imgs):
        if im.width != w:
            nh = int(im.height * w / im.width)
            im = im.resize((w, nh), Image.Resampling.LANCZOS)
        lab_h = 44 if labels else 0
        heights.append(im.height + lab_h)
        resized.append(im)
    total_h = sum(heights) + gap * (len(imgs) - 1) + 16
    canvas = Image.new("RGB", (w, total_h), (255, 255, 255))
    draw = ImageDraw.Draw(canvas)
    y = 8
    for i, im in enumerate(resized):
        if labels:
            draw.text((12, y), labels[i], font=font, fill=(15, 55, 95))
            y += 40
        canvas.paste(im, (0, y))
        y += im.height + gap
    canvas.save(dest, "PNG")


def sql_table(dest):
    import subprocess

    q = (
        "SET NOCOUNT ON; SELECT TOP 8 ID, CONVERT(varchar(19), Orderdate, 120) AS Orderdate, "
        "CustomerID, CAST(Totalprice AS bigint) AS Totalprice, "
        "CASE WHEN Status=1 THEN N'Da xu ly' ELSE N'Cho xu ly' END AS TrangThai "
        "FROM [Order] ORDER BY ID DESC;"
    )
    r = subprocess.run(
        ["sqlcmd", "-S", "localhost", "-E", "-d", "laptopstore", "-s", "|", "-W", "-Q", q],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    lines = [ln for ln in (r.stdout or "").splitlines() if ln.strip() and not ln.startswith("Msg")]
    font = ImageFont.truetype(r"C:\Windows\Fonts\consola.ttf", 22) if Path(r"C:\Windows\Fonts\consola.ttf").exists() else ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 20)
    title_f = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 26)
    body = "\n".join(lines[:20]) if lines else "(khong doc duoc SQL)"
    img = Image.new("RGB", (1400, 520), (255, 255, 255))
    d = ImageDraw.Draw(img)
    d.rectangle([8, 8, 1392, 512], outline=(21, 67, 120), width=2)
    d.text((24, 20), "SQL Server – catalog laptopstore – bang Order (SSMS / sqlcmd)", font=title_f, fill=(21, 67, 120))
    d.text((24, 70), body, font=font, fill=(30, 30, 30))
    img.save(dest, "PNG")


d1 = OUT / "HinhD1_login_register.png"
d2 = OUT / "HinhD2_timkiem_chitiet.png"
d4 = OUT / "HinhD4_crud_productnotsold.png"
d5 = OUT / "HinhD5_sql_order.png"
d3 = SHOT / "HinhD3_dathang.png"
h19 = SHOT / "Hinh19_admin_donhang.png"

vstack(
    [SHOT / "HinhD1_dangnhap.png", SHOT / "HinhD1_dangky.png"],
    d1,
    ["(a) Dang nhap khach – /Login/Login", "(b) Dang ky khach – /Login/Register"],
)
vstack(
    [SHOT / "HinhD2_timkiem_HP.png", SHOT / "HinhD2_chitiet.png"],
    d2,
    ["(a) Tim kiem Contains: HP", "(b) Chi tiet laptop HP 14 em0023AU"],
)
vstack(
    [SHOT / "HinhD4_items.png", SHOT / "HinhD4_productnotsold.png"],
    d4,
    ["(a) CRUD Items – /Items", "(b) Productnotsold – /Admin/Productnotsold"],
)
sql_table(d5)
shutil.copyfile(h18, OUT / "Hinh18_GiaoDienTrangChu.png")
shutil.copyfile(h19, OUT / "Hinh19_Admin.png")
shutil.copyfile(d3, OUT / "HinhD3_gio_dathang.png")


def has_drawing(p):
    xml = p._p.xml
    return "w:drawing" in xml or "v:imagedata" in xml or "w:pict" in xml


def insert_pic_before(caption_p, img_path, width_cm=15.5):
    prev = caption_p._p.getprevious()
    # xóa 1-2 đoạn ảnh ngay trước chú thích
    for _ in range(2):
        if prev is None:
            break
        if "w:drawing" in prev.xml or "v:imagedata" in prev.xml or "w:pict" in prev.xml:
            parent = prev.getparent()
            nxt = prev.getprevious()
            parent.remove(prev)
            prev = nxt
        else:
            break
    pic_elm = OxmlElement("w:p")
    caption_p._p.addprevious(pic_elm)
    pic_p = Paragraph(pic_elm, caption_p._parent)
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.add_run().add_picture(str(img_path), width=Cm(width_cm))


def set_caption(p, text):
    for r in p.runs:
        r.text = ""
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


doc = Document(str(DOC))
paras = list(doc.paragraphs)
seen18 = 0
seen19 = 0
for p in paras:
    t = (p.text or "").strip()
    if t.startswith("Hình 18. Giao diện trang chủ"):
        seen18 += 1
        insert_pic_before(p, OUT / "Hinh18_GiaoDienTrangChu.png")
        set_caption(p, "Hình 18. Giao diện trang chủ laptop.khannv.vn (localhost:51494)")
    elif t.startswith("Hình 19. Giao diện trang quản trị"):
        seen19 += 1
        insert_pic_before(p, OUT / "Hinh19_Admin.png")
        set_caption(p, "Hình 19. Giao diện trang quản trị Admin – danh sách đơn hàng")
    elif t.startswith("Hình D1."):
        insert_pic_before(p, d1, 15.2)
        set_caption(p, "Hình D1. Demo đăng nhập / đăng ký khách hàng (localhost:51494)")
    elif t.startswith("Hình D2."):
        insert_pic_before(p, d2, 15.2)
        set_caption(p, "Hình D2. Demo tìm kiếm Contains (HP) và trang chi tiết laptop")
    elif t.startswith("Hình D3."):
        insert_pic_before(p, OUT / "HinhD3_gio_dathang.png")
        set_caption(p, "Hình D3. Demo giỏ hàng / đặt hàng (Cart/Order)")
    elif t.startswith("Hình D4."):
        insert_pic_before(p, d4, 15.2)
        set_caption(p, "Hình D4. Demo CRUD Items và Productnotsold")
    elif t.startswith("Hình D5."):
        insert_pic_before(p, d5)
        set_caption(p, "Hình D5. Demo dữ liệu Order trên SQL Server catalog laptopstore")

# Xóa ảnh trang chủ mồ côi (không có chú thích Hình 18 ngay sau)
for i, p in enumerate(paras):
    t = (p.text or "").strip()
    if "Sản phẩm nỗi bật" in t or "Sản phẩm nổi bật tại http://localhost:51494" in t:
        nxt = paras[i + 1] if i + 1 < len(paras) else None
        if nxt is not None and has_drawing(nxt):
            n2 = paras[i + 2] if i + 2 < len(paras) else None
            n2t = (n2.text or "") if n2 else ""
            if not n2t.startswith("Hình 18"):
                nxt._p.getparent().remove(nxt._p)

doc.save(str(TMP))
shutil.copyfile(TMP, DOC)
TMP.unlink(missing_ok=True)
print("done 18", seen18, "19", seen19)
