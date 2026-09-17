# -*- coding: utf-8 -*-
"""Sửa 4 lỗi trước khi nộp: Chương 4, câu hướng dẫn, tên Hình 3–16."""
from pathlib import Path
from docx import Document

SRC = Path(r"C:\Users\nguye\OneDrive\Máy tính\Đồ án\BaoCao_laptop.khannv.vn.docx")


def set_text(p, text):
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""


def starts_num(t, prefix):
    return t.startswith(prefix)


def main():
    doc = Document(str(SRC))
    drop = []
    n_ch = n_fig = n_del = 0
    section = "pre"  # pre | ch3 | ch4install | ch4result | ch5

    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if not t:
            continue

        # --- delete instruction sentences ---
        low = t.lower()
        if (
            "cập nhật mục lục tự động" in low
            or "table of contents" in low
            or "khi in/bảo vệ" in low
            or "khi in báo cáo, sinh viên" in low
            or "sinh viên dán screenshot" in low
            or "sinh viên chèn ảnh" in low
            or "vẫn cần sinh viên chèn đầy đủ" in low
        ):
            drop.append(p._element)
            n_del += 1
            continue

        # track sections by headings
        if t.startswith("CHƯƠNG 3"):
            section = "ch3"
        elif t.startswith("3.4.") or t == "3.4. Cài đặt và triển khai chương trình" or t.startswith("3.4. Cài đặt"):
            section = "ch4install"
        elif t.startswith("CHƯƠNG 4") and "KẾT QUẢ" in t:
            section = "ch4result"
        elif t.startswith("CHƯƠNG 5") and "KẾT LUẬN" in t:
            section = "ch5"

        nt = t

        if section == "ch4install":
            if nt.startswith("3.4. Cài đặt"):
                nt = "CHƯƠNG 4. CÀI ĐẶT VÀ TRIỂN KHAI"
            elif nt.startswith("Mục 3.4 "):
                nt = "Chương 4 " + nt[len("Mục 3.4 ") :]
            elif nt.startswith("3.4."):
                nt = "4." + nt[4:]
        elif section == "ch4result":
            if nt.startswith("CHƯƠNG 4. KẾT QUẢ"):
                nt = "CHƯƠNG 5. KẾT QUẢ NGHIÊN CỨU"
            elif nt.startswith("Chương 4 tổng hợp"):
                nt = "Chương 5" + nt[len("Chương 4") :]
            elif nt.startswith("4."):
                nt = "5." + nt[2:]
        elif section == "ch5":
            if nt.startswith("CHƯƠNG 5. KẾT LUẬN"):
                nt = "KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN"

        # figure names: không gọi sơ đồ MVC là Class Diagram / biểu đồ lớp
        if "Biểu đồ lớp chức năng" in nt:
            nt = nt.replace("Biểu đồ lớp chức năng", "Sơ đồ xử lý chức năng")
            if "theo mô hình MVC" not in nt:
                # "Sơ đồ xử lý chức năng XXX (UCxx)" → thêm MVC trước (UC
                if "(UC" in nt:
                    nt = nt.replace(" (UC", " theo mô hình MVC (UC")
                else:
                    nt = nt.rstrip(".") + " theo mô hình MVC"
            n_fig += 1

        if "chỗ chèn biểu đồ lớp/lưu đồ" in nt:
            nt = nt.replace(" (kèm chỗ chèn biểu đồ lớp/lưu đồ)", "")
        if "Use Case, Class Diagram" in nt:
            nt = nt.replace(
                "Xây dựng biểu đồ UML (Use Case, Class Diagram) mô tả luồng hoạt động và cấu trúc hệ thống.",
                "Xây dựng biểu đồ Use Case và sơ đồ xử lý theo mô hình MVC mô tả luồng hoạt động của hệ thống.",
            )

        if nt != t:
            set_text(p, nt)
            if section in ("ch4install", "ch4result", "ch5") and (
                t.startswith("3.4") or t.startswith("4.") or t.startswith("CHƯƠNG 4") or t.startswith("CHƯƠNG 5") or t.startswith("Chương 4") or t.startswith("Mục 3.4")
            ):
                n_ch += 1

        # soften remaining screenshot paragraph (keep facts, no instruction)
        t2 = (p.text or "").strip()
        if t2.startswith("Các Hình 19–23 là ảnh chụp") or t2.startswith("Các Hình 19–23 minh họa"):
            set_text(
                p,
                "Các Hình 19–23 minh họa giao diện website chạy trên IIS Express (http://localhost:51494): trang chủ, trang quản trị, đăng nhập, giỏ hàng/đặt hàng, Productnotsold và dashboard doanh thu.",
            )

    for el in drop:
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)

    # tables
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    t = (p.text or "").strip()
                    if "Biểu đồ lớp" in t:
                        set_text(p, t.replace("Biểu đồ lớp chức năng", "Sơ đồ xử lý chức năng"))

    try:
        doc.save(str(SRC))
        print("saved", SRC)
    except PermissionError:
        out = SRC.with_name("_nop_fix.docx")
        doc.save(str(out))
        print("LOCKED", out)
    print("chapter", n_ch, "fig", n_fig, "deleted", n_del)


if __name__ == "__main__":
    main()
