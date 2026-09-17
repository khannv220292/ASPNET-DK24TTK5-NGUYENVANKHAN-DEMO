# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm
from docx.text.paragraph import Paragraph

docx = Path(r"C:\Users\nguye\OneDrive\Máy tính\Đồ án\BaoCao_laptop.khannv.vn_TheoMauChinh.docx")
img = Path(r"C:\Users\nguye\OneDrive\Máy tính\Đồ án\Hinh18_GiaoDienTrangChu.png")
caption_text = "Hình 18. Giao diện trang chủ laptop.khannv.vn (localhost:51494)"

doc = Document(str(docx))
target = None
for p in doc.paragraphs:
    t = (p.text or "").strip()
    if t.startswith("Hình 18") and "trang chủ" in t.lower():
        target = p
        break
if target is None:
    raise SystemExit("no hinh 18")

prev = target._p.getprevious()
if prev is not None and "w:drawing" in prev.xml:
    parent = prev.getparent()
    parent.remove(prev)

pic_elm = OxmlElement("w:p")
target._p.addprevious(pic_elm)
pic_p = Paragraph(pic_elm, target._parent)
pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
pic_p.add_run().add_picture(str(img), width=Cm(16.0))

for r in target.runs:
    r.text = ""
if target.runs:
    target.runs[0].text = caption_text
else:
    target.add_run(caption_text)
target.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.save(str(docx))
print("inserted")
