# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

src = Path(r"C:\Users\nguye\OneDrive\Máy tính\Đồ án\BaoCao_laptop.khannv.vn_TheoMauChinh.docx")
bak = Path(r"C:\Users\nguye\OneDrive\Máy tính\Đồ án\BaoCao_laptop.khannv.vn_TheoMauChinh.bak.docx")
img = Path(r"C:\Users\nguye\OneDrive\Máy tính\Đồ án\Hinh17_BieuDoLop_Productnotsold.png")

if not bak.exists():
    bak.write_bytes(src.read_bytes())

doc = Document(str(src))


def insert_paragraph_after(paragraph):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


target = None
for p in doc.paragraphs:
    t = p.text.strip()
    if "Hình 17" in t and "Productnotsold" in t and "Sinh viên chèn" in t:
        target = p
        break

if target is None:
    raise SystemExit("caption not found")

# Clean caption
for r in target.runs:
    r.text = ""
if target.runs:
    target.runs[0].text = "Hình 17. Biểu đồ lớp / màn hình Productnotsold (sản phẩm chưa bán)"
else:
    target.add_run("Hình 17. Biểu đồ lớp / màn hình Productnotsold (sản phẩm chưa bán)")
target.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Figure above caption (chuẩn báo cáo: ảnh rồi chú thích)
pic_elm = OxmlElement("w:p")
target._p.addprevious(pic_elm)
pic_p = Paragraph(pic_elm, target._parent)
pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = pic_p.add_run()
run.add_picture(str(img), width=Cm(16.0))

out = Path(r"C:\Users\nguye\OneDrive\Máy tính\Đồ án\BaoCao_laptop.khannv.vn_TheoMauChinh_Hinh17.docx")
doc.save(str(out))
print("saved", out)
