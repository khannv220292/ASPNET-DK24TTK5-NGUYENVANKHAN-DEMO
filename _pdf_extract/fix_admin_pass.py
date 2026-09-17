# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document

src = Path(r"C:\Users\nguye\OneDrive\Máy tính\Đồ án\BaoCao_laptop.khannv.vn_TheoMauChinh.docx")
out = Path(r"C:\Users\nguye\OneDrive\Máy tính\Đồ án\_tmp_baocao_admin.docx")
doc = Document(str(src))

repls = [
    ("Admin đăng nhập Admin/1.", "Admin đăng nhập Username: Admin, mật khẩu: admin."),
    ("tài khoản mẫu: Admin / 1", "tài khoản mẫu: Username Admin, mật khẩu admin"),
    ("Tài khoản demo: Admin / 1.", "Tài khoản demo: Username Admin, mật khẩu admin."),
    ("Đăng nhập Admin/1 →", "Đăng nhập Admin / mật khẩu admin →"),
    ("Mật khẩu mẫu Admin/1", "Mật khẩu mẫu admin"),
    ("Admin / 1 (bảng Admin", "Username Admin, mật khẩu admin (bảng Admin"),
    ("Nhớ tài khoản Admin/1;", "Nhớ tài khoản Admin / mật khẩu admin;"),
    ("tài khoản mẫu Admin/1 chỉ cho", "tài khoản mẫu Username Admin, mật khẩu admin chỉ cho"),
    ("đăng nhập Admin/1 để vào", "đăng nhập Username Admin, mật khẩu admin để vào"),
    ("Admin đăng nhập Admin/1,", "Admin đăng nhập Username Admin, mật khẩu admin,"),
    ("Admin/1", "Username Admin, mật khẩu admin"),
    ("Admin / 1", "Username Admin, mật khẩu admin"),
]

n = 0
for p in doc.paragraphs:
    t = p.text or ""
    nt = t
    for a, b in repls:
        nt = nt.replace(a, b)
    if nt != t:
        if p.runs:
            for r in p.runs:
                r.text = ""
            p.runs[0].text = nt
        else:
            p.add_run(nt)
        n += 1

for tbl in doc.tables:
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                t = p.text or ""
                nt = t
                for a, b in repls:
                    nt = nt.replace(a, b)
                if nt != t:
                    if p.runs:
                        for r in p.runs:
                            r.text = ""
                        p.runs[0].text = nt
                    else:
                        p.add_run(nt)
                    n += 1

doc.save(str(out))
print("paras", n)
