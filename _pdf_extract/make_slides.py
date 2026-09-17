# -*- coding: utf-8 -*-
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

OUT = Path(r"C:\Users\nguye\OneDrive\Máy tính\Đồ án") / "Slide_laptop.khannv.vn.pptx"
IMG = Path(r"C:\Users\nguye\OneDrive\Máy tính\Đồ án")

NAVY = RGBColor(0x15, 0x43, 0x78)
ACCENT = RGBColor(0x0B, 0x69, 0xB7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x55, 0x55, 0x55)
LIGHT = RGBColor(0xF4, 0xF7, 0xFB)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def bar(slide, y=0, h=0.12):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(y), prs.slide_width, Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = ACCENT
    sh.line.fill.background()


def footer(slide, n):
    bar(slide, 7.38, 0.12)
    tb = slide.shapes.add_textbox(Inches(0.4), Inches(7.15), Inches(10), Inches(0.28))
    p = tb.text_frame.paragraphs[0]
    p.text = "Nguyễn Văn Khan  |  DK24TTK5  |  Chuyên đề ASP.NET  |  laptop.khannv.vn"
    p.font.size = Pt(11)
    p.font.color.rgb = GRAY
    num = slide.shapes.add_textbox(Inches(12.2), Inches(7.15), Inches(0.9), Inches(0.28))
    np = num.text_frame.paragraphs[0]
    np.text = str(n)
    np.font.size = Pt(12)
    np.font.color.rgb = NAVY
    np.alignment = PP_ALIGN.RIGHT


def title_box(slide, text, top=0.28):
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(12.3), Inches(0.7))
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.font.name = "Calibri"


def add_bullets(tf, items, size=18):
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = it
        p.level = 0
        p.font.size = Pt(size)
        p.font.color.rgb = DARK
        p.font.name = "Calibri"
        p.space_after = Pt(8)


# ----- 1 Title -----
s = prs.slides.add_slide(prs.slide_layouts[6])
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = NAVY
bg.line.fill.background()
left = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.22), prs.slide_height)
left.fill.solid()
left.fill.fore_color.rgb = ACCENT
left.line.fill.background()
t = s.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.5), Inches(0.4))
p = t.text_frame.paragraphs[0]
p.text = "TRƯỜNG ĐẠI HỌC TRÀ VINH  •  CHUYÊN ĐỀ ASP.NET"
p.font.size = Pt(16)
p.font.color.rgb = RGBColor(0xC5, 0xD8, 0xF0)
t = s.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.8), Inches(1.8))
tf = t.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Xây dựng website bán laptop"
p.font.size = Pt(40)
p.font.bold = True
p.font.color.rgb = WHITE
p = tf.add_paragraph()
p.text = "laptop.khannv.vn  —  ASP.NET MVC 5, Entity Framework 6, SQL Server"
p.font.size = Pt(20)
p.font.color.rgb = RGBColor(0xD0, 0xE4, 0xF8)
t = s.shapes.add_textbox(Inches(0.8), Inches(4.6), Inches(11), Inches(1.6))
tf = t.text_frame
for line in [
    "Sinh viên: Nguyễn Văn Khan    MSSV: 170124589    Lớp: DK24TTK5",
    "Giảng viên hướng dẫn: TS. Nguyễn Nhứt Lam",
    "Demo: http://localhost:51494",
]:
    p = tf.paragraphs[0] if line.startswith("Sinh") else tf.add_paragraph()
    p.text = line
    p.font.size = Pt(18)
    p.font.color.rgb = WHITE

# ----- 2 Mục tiêu -----
s = prs.slides.add_slide(prs.slide_layouts[6])
bar(s, 0, 0.14)
title_box(s, "1. Đề tài và mục tiêu")
tb = s.shapes.add_textbox(Inches(0.6), Inches(1.15), Inches(12.1), Inches(5.6))
tf = tb.text_frame
tf.word_wrap = True
add_bullets(tf, [
    "Đề tài: website thương mại điện tử bán laptop (thương hiệu laptop.khannv.vn).",
    "Khách hàng: xem catalog, tìm kiếm gần đúng (LINQ Contains), giỏ Session, đặt hàng COD/MoMo.",
    "Quản trị: đăng nhập Admin, CRUD sản phẩm/danh mục, duyệt đơn, thống kê Productnotsold.",
    "Công nghệ: ASP.NET MVC 5, .NET Framework 4.8, EF6 (EDMX), SQL Server catalog laptopstore.",
    "Môi trường demo: Visual Studio, IIS Express cổng 51494.",
], 20)
footer(s, 2)

# ----- 3 MVC -----
s = prs.slides.add_slide(prs.slide_layouts[6])
bar(s, 0, 0.14)
title_box(s, "2. Mô hình ASP.NET MVC 5")
boxes = [
    (0.5, "Model", "Item, Customer, Order\nProTechTiveGearEntities\n(EF6 → SQL)"),
    (4.7, "View", "Razor .cshtml\n_LayoutHomePage\n_LayoutAdmin"),
    (8.9, "Controller", "Shop, Cart, Login\nAdmin, Items\nActionResult"),
]
for x, name, desc in boxes:
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.3), Inches(3.8), Inches(2.6))
    sh.fill.solid()
    sh.fill.fore_color.rgb = LIGHT
    sh.line.color.rgb = ACCENT
    tf = sh.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = name
    p.font.bold = True
    p.font.size = Pt(24)
    p.font.color.rgb = NAVY
    p.alignment = PP_ALIGN.CENTER
    p = tf.add_paragraph()
    p.text = desc
    p.font.size = Pt(16)
    p.font.color.rgb = DARK
    p.alignment = PP_ALIGN.CENTER
tb = s.shapes.add_textbox(Inches(0.6), Inches(4.2), Inches(12), Inches(2.5))
tf = tb.text_frame
tf.word_wrap = True
add_bullets(tf, [
    "Request: trình duyệt → IIS Express → RouteConfig {controller}/{action}/{id}.",
    "Mặc định: ShopController.Index (URL cũ /AuraStore vẫn nhận, không sinh link).",
    "Session[\"usr\"] khách, Session[\"Account\"] admin, Session[\"Cart\"] giỏ hàng.",
], 18)
footer(s, 3)

# ----- 4 CSDL -----
s = prs.slides.add_slide(prs.slide_layouts[6])
bar(s, 0, 0.14)
title_box(s, "3. Cơ sở dữ liệu SQL Server — catalog laptopstore")
tb = s.shapes.add_textbox(Inches(0.55), Inches(1.15), Inches(12.2), Inches(5.6))
tf = tb.text_frame
tf.word_wrap = True
add_bullets(tf, [
    "Tên CSDL: laptopstore (Web.config: Initial Catalog=laptopstore).",
    "Bảng chính: Admin, Customer, Menu (hãng), ItemType, Brand, Item, Order, OrderDetail, Payment, Banner, Blog.",
    "Quan hệ: Item thuộc Menu/ItemType; Order gắn Customer; OrderDetail gắn Item và Order.",
    "Script: Database/database.sql, seed_laptop.sql — import SSMS trước khi F5.",
    "ORM: Entity Framework 6, Gear.edmx, DbContext ProTechTiveGearEntities.",
], 19)
footer(s, 4)

# ----- 5 Chức năng khách -----
s = prs.slides.add_slide(prs.slide_layouts[6])
bar(s, 0, 0.14)
title_box(s, "4. Chức năng phía khách hàng")
tb = s.shapes.add_textbox(Inches(0.55), Inches(1.15), Inches(12.2), Inches(5.6))
tf = tb.text_frame
tf.word_wrap = True
add_bullets(tf, [
    "Trang chủ: banner, hãng nổi bật, sản phẩm mới, menu Lenovo / Dell / HP / ASUS…",
    "Tìm kiếm: Name, ShortTitle, Describe, loại, hãng — Contains → SQL LIKE.",
    "Chi tiết laptop, thêm giỏ (kiểm tra tồn kho), sửa/xóa giỏ.",
    "Đặt hàng: COD hoặc MoMo; cho phép khách vãng lai (guest theo SĐT).",
    "Đăng ký / đăng nhập / đổi mật khẩu (hash Encryption); theo dõi và hủy đơn.",
], 19)
footer(s, 5)

# ----- 6 Admin -----
s = prs.slides.add_slide(prs.slide_layouts[6])
bar(s, 0, 0.14)
title_box(s, "5. Chức năng phía quản trị")
tb = s.shapes.add_textbox(Inches(0.55), Inches(1.15), Inches(12.2), Inches(5.6))
tf = tb.text_frame
tf.word_wrap = True
add_bullets(tf, [
    "Đăng nhập: Username Admin  /  mật khẩu admin  → AllListOrder.",
    "CRUD Items, ItemType, Menu (hãng), Banner, Customer, Blog.",
    "Đơn hàng: tab tất cả / chưa xử lý / chưa thanh toán; xác nhận, xóa đơn.",
    "FeaturedBrand: chọn hãng nổi bật trên trang chủ.",
    "Báo cáo Productnotsold: laptop phục vụ tiêu chí thống kê học phần.",
], 19)
footer(s, 6)

# ----- 7 Ảnh trang chủ -----
s = prs.slides.add_slide(prs.slide_layouts[6])
bar(s, 0, 0.14)
title_box(s, "6. Giao diện trang chủ (demo)")
pic = IMG / "Hinh18_GiaoDienTrangChu.png"
if pic.exists():
    s.shapes.add_picture(str(pic), Inches(1.4), Inches(1.15), width=Inches(10.5))
footer(s, 7)

# ----- 8 Tìm kiếm -----
s = prs.slides.add_slide(prs.slide_layouts[6])
bar(s, 0, 0.14)
title_box(s, "7. Tìm kiếm Contains và chi tiết sản phẩm")
pic = IMG / "HinhD2_timkiem_chitiet.png"
if pic.exists():
    s.shapes.add_picture(str(pic), Inches(2.4), Inches(1.05), height=Inches(5.9))
footer(s, 8)

# ----- 9 Admin screenshot -----
s = prs.slides.add_slide(prs.slide_layouts[6])
bar(s, 0, 0.14)
title_box(s, "8. Giao diện quản trị — danh sách đơn hàng")
pic = IMG / "Hinh19_Admin.png"
if pic.exists():
    s.shapes.add_picture(str(pic), Inches(1.2), Inches(1.15), width=Inches(10.9))
footer(s, 9)

# ----- 10 Kết luận -----
s = prs.slides.add_slide(prs.slide_layouts[6])
bar(s, 0, 0.14)
title_box(s, "9. Kết luận và hướng phát triển")
tb = s.shapes.add_textbox(Inches(0.55), Inches(1.1), Inches(12.2), Inches(5.7))
tf = tb.text_frame
tf.word_wrap = True
add_bullets(tf, [
    "Đã hoàn thành website bán laptop trên ASP.NET MVC 5 + SQL Server laptopstore.",
    "Đáp ứng luồng bán hàng, CRUD quản trị, tìm kiếm tương đối, báo cáo Productnotsold.",
    "Hạn chế: cổng thanh toán ngân hàng thật, ASP.NET Identity đầy đủ, chưa lên ASP.NET Core.",
    "Hướng mở: lọc cấu hình, phân trang, HTTPS, Identity, báo cáo doanh thu theo ngày.",
    "GitHub: github.com/khannv220292/ASPNET-DK24TTK5-NGUYENVANKHAN-DEMO",
    "Demo tài khoản Admin / admin  •  F5 → http://localhost:51494",
], 18)
footer(s, 10)

# ----- 11 Cảm ơn -----
s = prs.slides.add_slide(prs.slide_layouts[6])
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = NAVY
bg.line.fill.background()
t = s.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.7), Inches(2.2))
tf = t.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Xin cảm ơn quý thầy cô"
p.font.size = Pt(40)
p.font.bold = True
p.font.color.rgb = WHITE
p.alignment = PP_ALIGN.CENTER
p = tf.add_paragraph()
p.text = "Sẵn sàng demo website và trả lời câu hỏi"
p.font.size = Pt(20)
p.font.color.rgb = RGBColor(0xC5, 0xD8, 0xF0)
p.alignment = PP_ALIGN.CENTER

prs.save(str(OUT))
print("saved", OUT, OUT.stat().st_size)
