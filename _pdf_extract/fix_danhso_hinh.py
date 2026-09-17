# -*- coding: utf-8 -*-
"""Sửa đánh số chương, thống nhất tên hình, bổ sung mô tả thống kê."""
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

SRC = Path(r"C:\Users\nguye\OneDrive\Máy tính\Đồ án\BaoCao_laptop.khannv.vn.docx")

# Thứ tự: khóa dài trước
CHAPTER = [
    ("4.3.4.", "3.4.3.4."),
    ("4.3.3.", "3.4.3.3."),
    ("4.3.2.", "3.4.3.2."),
    ("4.3.1.", "3.4.3.1."),
    ("4.3.", "3.4.3."),
    ("4.2.6.", "3.4.2.6."),
    ("4.2.5.", "3.4.2.5."),
    ("4.2.4.", "3.4.2.4."),
    ("4.2.3.", "3.4.2.3."),
    ("4.2.2.", "3.4.2.2."),
    ("4.2.1.", "3.4.2.1."),
    ("4.2.", "3.4.2."),
    ("4.1.", "3.4.1."),
    ("5.4.2.", "4.4.2."),
    ("5.4.1.", "4.4.1."),
    ("5.4.", "4.4."),
    ("5.3.", "4.3."),
    ("5.2.3.", "4.2.3."),
    ("5.2.2.", "4.2.2."),
    ("5.2.1.", "4.2.1."),
    ("5.2.", "4.2."),
    ("5.1.3.", "4.1.3."),
    ("5.1.2.", "4.1.2."),
    ("5.1.1.", "4.1.1."),
    ("5.1.", "4.1."),
]

# Thống nhất chú thích theo thứ tự xuất hiện trong thân bài
CAPTION_EXACT = {
    "Hình 1. Mô hình ASP.NET MVC": "Hình 1. Mô hình ASP.NET MVC",
    "Hình 20. Sơ đồ cơ sở dữ liệu SQL Server – ERD catalog laptopstore": "Hình 2. Sơ đồ cơ sở dữ liệu SQL Server (ERD, catalog laptopstore)",
    "Hình 20. Sơ đồ cơ sở dữ liệu SQL Server (laptopstore)": "Hình 2. Sơ đồ cơ sở dữ liệu SQL Server (ERD, catalog laptopstore)",
    "Hình 2. Mô hình use case của các tác nhân (Khách hàng – Quản trị)": "Hình 3. Biểu đồ Use Case hệ thống website bán laptop",
    "Hình 2. Biểu đồ Use Case hệ thống website bán laptop": "Hình 3. Biểu đồ Use Case hệ thống website bán laptop",
    "Hình 3. Biểu đồ lớp chức năng đăng ký": "Hình 4. Biểu đồ lớp chức năng đăng ký tài khoản (UC01)",
    "Hình 3. Biểu đồ lớp chức năng đăng ký tài khoản (UC01)": "Hình 4. Biểu đồ lớp chức năng đăng ký tài khoản (UC01)",
    "Hình 4. Biểu đồ lớp chức năng đăng nhập": "Hình 5. Biểu đồ lớp chức năng đăng nhập (UC02)",
    "Hình 4. Biểu đồ lớp chức năng đăng nhập (UC02)": "Hình 5. Biểu đồ lớp chức năng đăng nhập (UC02)",
    "Hình 5. Biểu đồ lớp chức năng cập nhật thông tin / đổi mật khẩu": "Hình 6. Biểu đồ lớp chức năng cập nhật hồ sơ / đổi mật khẩu (UC03)",
    "Hình 5. Biểu đồ lớp chức năng cập nhật hồ sơ / đổi mật khẩu (UC03)": "Hình 6. Biểu đồ lớp chức năng cập nhật hồ sơ / đổi mật khẩu (UC03)",
    "Hình 6. Biểu đồ lớp chức năng xem chi tiết sản phẩm": "Hình 7. Biểu đồ lớp chức năng xem chi tiết laptop (UC04)",
    "Hình 6. Biểu đồ lớp chức năng xem chi tiết laptop (UC04)": "Hình 7. Biểu đồ lớp chức năng xem chi tiết laptop (UC04)",
    "Hình 7. Biểu đồ lớp chức năng xem danh sách / trang chủ sản phẩm": "Hình 8. Biểu đồ lớp chức năng trang chủ / danh sách sản phẩm (UC05)",
    "Hình 7. Biểu đồ lớp chức năng trang chủ / danh sách sản phẩm (UC05)": "Hình 8. Biểu đồ lớp chức năng trang chủ / danh sách sản phẩm (UC05)",
    "Hình 8. Biểu đồ lớp chức năng tìm kiếm sản phẩm (Contains)": "Hình 9. Biểu đồ lớp chức năng tìm kiếm tương đối Contains (UC06)",
    "Hình 8. Biểu đồ lớp chức năng tìm kiếm tương đối Contains (UC06)": "Hình 9. Biểu đồ lớp chức năng tìm kiếm tương đối Contains (UC06)",
    "Hình 9. Biểu đồ lớp chức năng theo dõi đơn hàng": "Hình 10. Biểu đồ lớp chức năng theo dõi đơn hàng của khách (UC07)",
    "Hình 9. Biểu đồ lớp chức năng theo dõi đơn hàng của khách (UC07)": "Hình 10. Biểu đồ lớp chức năng theo dõi đơn hàng của khách (UC07)",
    "Hình 10. Biểu đồ lớp chức năng hủy / cập nhật trạng thái đơn hàng": "Hình 11. Biểu đồ lớp chức năng hủy / cập nhật trạng thái đơn (UC08)",
    "Hình 10. Biểu đồ lớp chức năng hủy đơn hàng": "Hình 11. Biểu đồ lớp chức năng hủy / cập nhật trạng thái đơn (UC08)",
    "Hình 11. Biểu đồ lớp chức năng quản lý giỏ hàng (Session)": "Hình 12. Biểu đồ lớp chức năng quản lý giỏ hàng Session (UC09)",
    "Hình 11. Biểu đồ lớp chức năng quản lý giỏ hàng Session (UC09)": "Hình 12. Biểu đồ lớp chức năng quản lý giỏ hàng Session (UC09)",
    "Hình 12. Biểu đồ lớp chức năng đặt hàng": "Hình 13. Biểu đồ lớp chức năng đặt hàng (UC10)",
    "Hình 12. Biểu đồ lớp chức năng đặt hàng (UC10)": "Hình 13. Biểu đồ lớp chức năng đặt hàng (UC10)",
    "Hình 13. Biểu đồ lớp chức năng quản lý sản phẩm (Items CRUD)": "Hình 14. Biểu đồ lớp chức năng quản lý laptop Items CRUD (UC11)",
    "Hình 13. Biểu đồ lớp chức năng quản lý laptop Items CRUD (UC11)": "Hình 14. Biểu đồ lớp chức năng quản lý laptop Items CRUD (UC11)",
    "Hình 14. Biểu đồ lớp chức năng quản lý Brand / ItemType / Menu": "Hình 15. Biểu đồ lớp chức năng Brand / ItemType / Menu (UC12)",
    "Hình 14. Biểu đồ lớp chức năng Brand / ItemType / Menu (UC12)": "Hình 15. Biểu đồ lớp chức năng Brand / ItemType / Menu (UC12)",
    "Hình 15. Biểu đồ lớp chức năng quản lý Customer": "Hình 16. Biểu đồ lớp chức năng quản lý khách hàng (UC13)",
    "Hình 15. Biểu đồ lớp chức năng quản lý khách hàng (UC13)": "Hình 16. Biểu đồ lớp chức năng quản lý khách hàng (UC13)",
    "Hình 16. Biểu đồ lớp chức năng quản lý đơn hàng Admin": "Hình 17. Biểu đồ lớp chức năng quản lý đơn hàng Admin (UC14)",
    "Hình 16. Biểu đồ lớp chức năng quản lý đơn hàng Admin (UC14)": "Hình 17. Biểu đồ lớp chức năng quản lý đơn hàng Admin (UC14)",
    "Hình 17. Sơ đồ xử lý chức năng Productnotsold theo mô hình MVC": "Hình 18. Sơ đồ xử lý chức năng Productnotsold theo mô hình MVC (UC15)",
    "Hình 17. Biểu đồ lớp / màn hình Productnotsold (sản phẩm chưa bán)": "Hình 18. Sơ đồ xử lý chức năng Productnotsold theo mô hình MVC (UC15)",
    "Hình 18. Giao diện trang chủ laptop.khannv.vn (localhost:51494)": "Hình 19. Giao diện trang chủ laptop.khannv.vn (localhost:51494)",
    "Hình 18. Giao diện trang chủ laptop.khannv.vn": "Hình 19. Giao diện trang chủ laptop.khannv.vn (localhost:51494)",
    "Hình 19. Giao diện trang quản trị Admin": "Hình 20. Giao diện trang quản trị Admin (localhost:51494)",
    "Hình D1. Demo đăng nhập": "Hình 21. Minh chứng giao diện đăng nhập khách hàng",
    "Hình D2. Demo giỏ hàng Session và đặt hàng": "Hình 22. Minh chứng giỏ hàng Session và đặt hàng",
    "Hình D3. Demo CRUD Items và Productnotsold": "Hình 23. Minh chứng CRUD Items, Productnotsold và báo cáo doanh thu",
}

STATS = (
    "Báo cáo/thống kê gồm hai phần. (1) Productnotsold: AdminController.Productnotsold "
    "lấy Item không xuất hiện trong OrderDetail (LINQ Distinct ItemId, Where !Contains), "
    "View Productnotsold.cshtml liệt kê laptop chưa bán. (2) Dashboard doanh thu: AdminController.Index "
    "tổng hợp Order.Totalprice theo 12 tháng (lọc Month và Year), có lọc khoảng ngày trên giao diện; "
    "hiển thị KPI số đơn, đã thu/chưa thu, số lượng bán — phục vụ tiêu chí báo cáo trên phiếu đánh giá."
)

CODE_STATS = (
    "var soldIds = db.OrderDetails.Select(od => od.ItemId).Distinct();\n"
    "var notSold = db.Items.Where(i => !soldIds.Contains(i.ID)).ToList();\n"
    "return View(notSold);"
)


def set_para(p, text):
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""


def replace_prefix(text, pairs):
    for a, b in pairs:
        if text.startswith(a):
            return b + text[len(a):]
    return text


def main():
    doc = Document(str(SRC))
    n_ch = n_fig = 0
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if not t:
            continue
        nt = replace_prefix(t, CHAPTER)
        if nt != t:
            set_para(p, nt)
            t = nt
            n_ch += 1
        if t in CAPTION_EXACT and CAPTION_EXACT[t] != t:
            set_para(p, CAPTION_EXACT[t])
            n_fig += 1
            t = CAPTION_EXACT[t]
        elif t.startswith("Hình ") and t in CAPTION_EXACT:
            pass
        # intro 3.4
        if t.startswith("Chương này mô tả cấu trúc dự án"):
            set_para(p, "Mục 3.4 mô tả cấu trúc dự án, hiện thực các module chức năng chính kèm đoạn mã minh họa, và giao diện chạy thực tế trên IIS Express.")
        if t.startswith("Chương này tổng hợp kết quả"):
            set_para(p, "Chương 4 tổng hợp kết quả đạt được của đồ án website bán laptop laptop.khannv.vn, minh họa bằng ảnh màn hình chạy thật và đối chiếu mục tiêu.")

        # old productnotsold code block
        if "from t1 in db.Items" in t or "Laptop chưa từng xuất hiện" in t:
            set_para(p, CODE_STATS)
        if t.startswith("Module báo cáo phục vụ quản trị xem laptop chưa"):
            set_para(p, STATS)

    # tables text
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    t = (p.text or "").strip()
                    nt = replace_prefix(t, CHAPTER)
                    if nt != t:
                        set_para(p, nt)
                    if t in CAPTION_EXACT:
                        set_para(p, CAPTION_EXACT[t])
                    if "Productnotsold" in t and "Đạt" in t or (t.startswith("Báo cáo/thống kê")):
                        if "chỉ Productnotsold" in t or t.startswith("Báo cáo/thống kê —"):
                            set_para(
                                p,
                                "Báo cáo/thống kê — hiện thực (Productnotsold: Item chưa có trong OrderDetail; dashboard doanh thu 12 tháng trên Admin/Index).",
                            )

    try:
        doc.save(str(SRC))
        print("saved", SRC)
    except PermissionError:
        alt = SRC.with_name("BaoCao_laptop.khannv.vn.docx")
        out = SRC.with_name("_fix_danhso.docx")
        doc.save(str(out))
        print("LOCKED", out)
    print("chapter", n_ch, "fig", n_fig)


if __name__ == "__main__":
    main()
